"""Importa a planilha-semente (curadoria manual de 24/09/2026).

Três abas com as mesmas colunas: Editais, Lei do Esporte, Parecerista.
Colunas: # | Edital / Programa | Tipo de apoio | Abrangência | Prazo | Data-limite |
Dias restantes | Resumo | Link de inscrição | Observações | Fonte | Novo (24/09)

"Dias restantes" e "Novo" não são importados: o sistema calcula.
Linhas sem número na coluna "#" são notas de rodapé e são ignoradas.

Camadas, nesta ordem:
1. regras automáticas (config/mapeamentos.yaml);
2. data/seed/ajustes_seed.yaml — transcrição de informações que estão no texto
   da própria linha mas que as regras não conseguem ler (órgão, lista de cidades...).
   A "nota" de um ajuste é acrescentada às observações;
3. dedupe entre abas;
4. data/overrides.yaml (curadoria por id) — aplicado depois, no pipeline.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from pipeline.config import SEED_AJUSTES, SEED_XLSX, ler_yaml
from pipeline.dedupe import deduplicar
from pipeline.enums import TipoApoio
from pipeline.ids import gerar_id
from pipeline.models import Edital, FonteSecundaria
from pipeline.normalize import regras
from pipeline.textnorm import normalizar, url_generica

FONTE_SEED = "curadoria inicial (planilha 24/09/2026)"
CAPTURADO_EM = datetime(2026, 9, 24, 12, 0, tzinfo=ZoneInfo("America/Bahia"))
ABAS = {"Editais": "editais", "Lei do Esporte": "lei_do_esporte", "Parecerista": "parecerista"}
COLUNAS = [
    "#", "Edital / Programa", "Tipo de apoio", "Abrangência", "Prazo", "Data-limite",
    "Dias restantes", "Resumo", "Link de inscrição", "Observações", "Fonte",
]


@dataclass
class ResultadoImportacao:
    linhas_lidas: int = 0
    editais: list[Edital] = field(default_factory=list)
    fusoes: list[tuple[str, str]] = field(default_factory=list)
    pendencias: list[tuple[str, str]] = field(default_factory=list)  # (id, motivo)


def _texto(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _data(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        d, m, a = v.strip()[:10].split("/")
        return date(int(a), int(m), int(d))
    return None


def ler_linhas(caminho: Path) -> list[dict[str, Any]]:
    wb = load_workbook(caminho, data_only=False, read_only=True)
    linhas: list[dict[str, Any]] = []
    for nome_aba, chave_aba in ABAS.items():
        ws = wb[nome_aba]
        rows = list(ws.iter_rows(values_only=True))
        cabecalho = [str(c).strip() if c else "" for c in rows[0]]
        if cabecalho[: len(COLUNAS)] != COLUNAS:
            raise ValueError(f"Aba {nome_aba}: colunas inesperadas {cabecalho}")
        for r in rows[1:]:
            if len(r) < 2 or r[0] is None or not _texto(r[1]):
                continue  # nota de rodapé ou linha vazia
            d = dict(zip(COLUNAS, r, strict=False))
            d["_aba"] = chave_aba
            d["_ref"] = f"{nome_aba}#{int(str(r[0]).split('.')[0])}"
            linhas.append(d)
    wb.close()
    return linhas


def linha_para_edital(d: dict[str, Any], ajuste: dict[str, Any]) -> tuple[Edital, list[str]]:
    titulo = _texto(d["Edital / Programa"]) or ""
    tipo_txt = _texto(d["Tipo de apoio"])
    abr_txt = _texto(d["Abrangência"])
    prazo = _texto(d["Prazo"])
    resumo = _texto(d["Resumo"])
    obs = _texto(d["Observações"])
    link = _texto(d["Link de inscrição"])
    pendencias: list[str] = []

    tipo = regras.tipo_apoio(tipo_txt)
    abrangencia, ufs, cidades = regras.abrangencia(abr_txt)
    textos_mec = (titulo, tipo_txt, abr_txt, resumo)
    link_edital = regras.link_edital_das_observacoes(obs)
    data_limite = _data(d["Data-limite"])
    # "Dias restantes" é uma fórmula que devolve "Fluxo contínuo" para qualquer linha
    # sem data; por isso só o texto do prazo decide.
    fluxo = "fluxo continuo" in normalizar(prazo)

    marcadores = ["carga_inicial", f"aba:{d['_aba']}"]
    areas = regras.areas(titulo, tipo_txt, resumo)
    if d["_aba"] == "lei_do_esporte" and "esporte" not in areas:
        areas.append("esporte")

    campos: dict[str, Any] = {
        "id": gerar_id(FONTE_SEED, titulo),
        "titulo": titulo,
        "orgao": None,
        "tipo_apoio": tipo,
        "mecanismo": regras.mecanismos(*textos_mec),
        "leis_estaduais": regras.leis_estaduais(*textos_mec),
        "exige_projeto_aprovado": regras.exige_projeto_aprovado(
            tipo, titulo, tipo_txt, resumo, obs
        ),
        "abrangencia": abrangencia,
        "ufs": ufs,
        "cidades": cidades,
        "areas": areas,
        "publico_prioritario": regras.publico_prioritario(titulo, resumo, abr_txt),
        "proponente": regras.proponentes(tipo_txt, abr_txt, resumo),
        **regras.valores(resumo),
        "prazo_texto": prazo,
        "data_limite": data_limite,
        "fluxo_continuo": fluxo,
        "resumo": resumo,
        "link_inscricao": link,
        "link_edital": link_edital,
        "observacoes": obs,
        "fonte": FONTE_SEED,
        "source_url": link_edital or link,
        "fontes_secundarias": [
            FonteSecundaria(fonte=f"{_texto(d['Fonte'])} (via planilha, {d['_ref']})")
        ] if _texto(d["Fonte"]) else [],
        "captured_at": CAPTURADO_EM,
        "updated_at": CAPTURADO_EM,
        "marcadores": marcadores,
    }

    # Camada 2: ajustes transcritos da própria linha.
    nota = ajuste.pop("nota", None)
    for k, v in ajuste.items():
        if k not in campos:
            raise KeyError(f"{d['_ref']}: campo desconhecido em ajustes_seed.yaml: {k}")
        campos[k] = v
    if nota:
        campos["observacoes"] = " | ".join(x for x in (campos["observacoes"], nota) if x)

    if url_generica(link):
        marcadores.append("link_generico")
        pendencias.append("link de inscrição genérico (listagem da Prosas)")
    if regras.dominio_nao_oficial(link):
        marcadores.append("fonte_nao_oficial")
        pendencias.append("link aponta para notícia/agregador, não para a fonte oficial")
    for nome in ("orgao", "tipo_apoio", "abrangencia"):
        if campos[nome] in (None, ""):
            pendencias.append(f"{nome} não identificado")
    if not campos["mecanismo"] and campos["tipo_apoio"] in (
        TipoApoio.patrocinio_incentivado, TipoApoio.lei_incentivo_janela,
    ):
        pendencias.append("mecanismo não identificado")

    campos["revisao_pendente"] = bool(pendencias)
    return Edital(**campos), pendencias


def importar(caminho: Path = SEED_XLSX, ajustes_path: Path = SEED_AJUSTES) -> ResultadoImportacao:
    ajustes: dict[str, dict[str, Any]] = ler_yaml(ajustes_path) or {}
    res = ResultadoImportacao()
    linhas = ler_linhas(caminho)
    res.linhas_lidas = len(linhas)

    brutos: list[Edital] = []
    pend: dict[str, list[str]] = {}
    for d in linhas:
        e, p = linha_para_edital(d, dict(ajustes.get(d["_ref"], {})))
        brutos.append(e)
        if p:
            pend[e.id] = p

    res.editais, res.fusoes = deduplicar(brutos)
    ids_finais = {e.id: e for e in res.editais}
    for id_, motivos in pend.items():
        destino = id_ if id_ in ids_finais else next(
            e.id for e in res.editais if id_ in e.aliases
        )
        res.pendencias += [(destino, m) for m in motivos]
    return res

