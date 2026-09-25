"""Exportações: JSON (fonte da verdade + API), CSV e XLSX."""

import csv
import io
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.models import SCHEMA_VERSION, Edital

LICENCA = "CC-BY-4.0"
ATRIBUICAO = (
    "Central de Editais — Flores Produções e Ponto de Cultura Lab Garra. "
    "Dados sob CC-BY 4.0; cite também a fonte original de cada registro (campo source_url)."
)

# Colunas das planilhas (CSV/XLSX), na ordem de leitura humana.
COLUNAS_PLANILHA: list[tuple[str, str]] = [
    ("id", "ID"),
    ("titulo", "Edital / Programa"),
    ("orgao", "Órgão"),
    ("status", "Status"),
    ("prioridade", "Prioridade"),
    ("tipo_apoio", "Tipo de apoio"),
    ("mecanismo", "Mecanismo"),
    ("leis_estaduais", "Leis estaduais"),
    ("exige_projeto_aprovado", "Exige projeto aprovado"),
    ("abrangencia", "Abrangência"),
    ("ufs", "UFs"),
    ("cidades", "Cidades"),
    ("areas", "Áreas"),
    ("publico_prioritario", "Público prioritário"),
    ("proponente", "Proponente"),
    ("valor_por_projeto_min", "Valor mín. por projeto (R$)"),
    ("valor_por_projeto_max", "Valor máx. por projeto (R$)"),
    ("valor_total", "Valor total (R$)"),
    ("valor_texto", "Valores (texto da fonte)"),
    ("prazo_texto", "Prazo"),
    ("data_limite", "Data-limite"),
    ("fluxo_continuo", "Fluxo contínuo"),
    ("data_abertura", "Abertura"),
    ("resumo", "Resumo"),
    ("link_inscricao", "Link de inscrição"),
    ("link_edital", "Edital oficial"),
    ("observacoes", "Observações"),
    ("fonte", "Fonte"),
    ("source_url", "URL da fonte"),
    ("captured_at", "Capturado em"),
    ("updated_at", "Atualizado em"),
]


def ordenar(editais: list[Edital]) -> list[Edital]:
    return sorted(editais, key=lambda e: e.id)


def documento_json(editais: list[Edital], gerado_em: datetime | None = None) -> dict[str, Any]:
    """gerado_em só entra nas cópias publicadas (API). No data/editais.json ele fica
    de fora para o arquivo só mudar quando os dados mudam (diff limpo no git)."""
    doc: dict[str, Any] = {"schema": SCHEMA_VERSION}
    if gerado_em is not None:
        doc["gerado_em"] = gerado_em.isoformat(timespec="seconds")
    doc |= {
        "licenca": LICENCA,
        "atribuicao": ATRIBUICAO,
        "total": len(editais),
        "editais": [e.para_json() for e in ordenar(editais)],
    }
    return doc


def escrever_json(editais: list[Edital], destino: Path, gerado_em: datetime | None = None) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    texto = json.dumps(documento_json(editais, gerado_em), ensure_ascii=False, indent=2)
    destino.write_text(texto + "\n", encoding="utf-8")


def ler_json(origem: Path) -> list[Edital]:
    dados = json.loads(origem.read_text(encoding="utf-8"))
    return [Edital.model_validate(d) for d in dados["editais"]]


def _celula(valor: Any) -> Any:
    if isinstance(valor, list):
        return "; ".join(str(v) for v in valor)
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    return valor


def linhas_planilha(editais: list[Edital]) -> list[list[Any]]:
    linhas = []
    for e in ordenar(editais):
        d = e.para_json()
        linhas.append([_celula(d.get(campo)) for campo, _ in COLUNAS_PLANILHA])
    return linhas


def escrever_csv(editais: list[Edital], destino: Path) -> None:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([rotulo for _, rotulo in COLUNAS_PLANILHA])
    w.writerows(["" if v is None else v for v in linha] for linha in linhas_planilha(editais))
    destino.parent.mkdir(parents=True, exist_ok=True)
    # BOM para o Excel reconhecer UTF-8 (acentos) ao abrir com dois cliques.
    destino.write_text("﻿" + buf.getvalue(), encoding="utf-8")


def escrever_xlsx(editais: list[Edital], destino: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Editais"
    ws.append([rotulo for _, rotulo in COLUNAS_PLANILHA])
    for linha in linhas_planilha(editais):
        ws.append(linha)

    cabecalho = PatternFill("solid", fgColor="1F3A2E")
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF", size=12)
        cel.fill = cabecalho
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    larguras = {"titulo": 50, "resumo": 80, "observacoes": 50, "orgao": 28}
    for i, (campo, _) in enumerate(COLUNAS_PLANILHA, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(campo, 18)
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    sobre = wb.create_sheet("Sobre")
    sobre.append(["Central de Editais"])
    sobre.append([ATRIBUICAO])
    sobre.append(["Licença dos dados", LICENCA])
    sobre.append(["Gerado em", datetime.now(UTC).isoformat(timespec="seconds")])
    sobre.column_dimensions["A"].width = 100

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
