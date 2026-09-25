"""Usa a planilha-semente versionada em data/seed/ — nenhum acesso à internet."""

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pipeline.enums import Abrangencia, Mecanismo
from pipeline.export import formatos
from pipeline.importers.seed_xlsx import ResultadoImportacao, importar


@pytest.fixture(scope="module")
def seed() -> ResultadoImportacao:
    return importar()


def _por_titulo(res: ResultadoImportacao, inicio: str):  # type: ignore[no-untyped-def]
    return next(e for e in res.editais if e.titulo.startswith(inicio))


def test_contagens(seed: ResultadoImportacao) -> None:
    assert seed.linhas_lidas == 63
    assert len(seed.fusoes) == 7
    assert len(seed.editais) == 56


def test_todo_registro_tem_fonte_orgao_e_tipo(seed: ResultadoImportacao) -> None:
    for e in seed.editais:
        assert str(e.source_url).startswith("http"), e.id
        assert e.captured_at is not None
        assert e.orgao, e.id
        assert e.tipo_apoio is not None, e.id


def test_ids_unicos(seed: ResultadoImportacao) -> None:
    ids = [e.id for e in seed.editais]
    assert len(ids) == len(set(ids))


def test_fusao_entre_abas_une_mecanismos(seed: ResultadoImportacao) -> None:
    mapfre = _por_titulo(seed, "Mapfre")
    assert {Mecanismo.rouanet, Mecanismo.lei_esporte} <= set(mapfre.mecanismo)
    assert {"aba:editais", "aba:lei_do_esporte"} <= set(mapfre.marcadores)


def test_editais_diferentes_nao_se_fundem(seed: ResultadoImportacao) -> None:
    titulos = {e.titulo for e in seed.editais}
    assert "Clipe da Quebrada 2026 (APAA)" in titulos
    assert any(t.startswith("Rouanet nas Favelas") for t in titulos)
    assert any(t.startswith("Lei Rouanet 2026") for t in titulos)


def test_ajustes_aplicados(seed: ResultadoImportacao) -> None:
    ambev = _por_titulo(seed, "Ambev")
    assert ambev.abrangencia == Abrangencia.nacional and ambev.cidades == []
    assert ambev.exige_projeto_aprovado is True
    assert "Fazcultura/BA" in ambev.leis_estaduais
    favelas = _por_titulo(seed, "Rouanet nas Favelas")
    assert "Recife" in favelas.cidades and favelas.valor_por_projeto_max == 200_000


def test_links_genericos_vao_para_revisao(seed: ResultadoImportacao) -> None:
    genericos = [e for e in seed.editais if "link_generico" in e.marcadores]
    assert len(genericos) == 4 and all(e.revisao_pendente for e in genericos)


def test_exportacoes(seed: ResultadoImportacao, tmp_path: Path) -> None:
    formatos.escrever_json(seed.editais, tmp_path / "e.json")
    doc = json.loads((tmp_path / "e.json").read_text(encoding="utf-8"))
    assert doc["total"] == 56 and doc["licenca"] == "CC-BY-4.0" and "gerado_em" not in doc
    assert [e.id for e in formatos.ler_json(tmp_path / "e.json")] == sorted(
        e.id for e in seed.editais
    )

    formatos.escrever_csv(seed.editais, tmp_path / "e.csv")
    bruto = (tmp_path / "e.csv").read_text(encoding="utf-8")
    assert bruto.startswith("﻿")
    linhas = list(csv.reader(bruto.lstrip("﻿").splitlines()))
    assert linhas[0][1] == "Edital / Programa" and len(linhas) == 57

    formatos.escrever_xlsx(seed.editais, tmp_path / "e.xlsx")
    wb = load_workbook(tmp_path / "e.xlsx")
    assert wb.sheetnames == ["Editais", "Sobre"] and wb["Editais"].max_row == 57
